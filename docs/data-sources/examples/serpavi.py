# /// script
# requires-python = ">=3.11"
# dependencies = ["httpx", "openpyxl"]
# ///
"""Download the SERPAVI database (Ministerio de Vivienda) and look up the market rent range of a municipality.

Usage:
    uv run docs/data-sources/examples/serpavi.py                 # Madrid, 70 m2
    uv run docs/data-sources/examples/serpavi.py Valencia 85

The first run downloads ~71 MB into data/raw/ (git-ignored); the next runs reuse it.
"""

import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import httpx
from openpyxl import load_workbook

URL = (
    "https://cdn.mivau.gob.es/portal-web-mivau/vivienda/serpavi/"
    "2026-03_09_bd_SERPAVI_2011-2024%20-%20DEFINITIVO%20WEB_v2.xlsx"
)
FILE = Path("data/raw/serpavi_2011-2024.xlsx")
YEAR = "24"  # last year available in this edition (2024)
# The ministry site answers 403 to requests without a browser User-Agent.
HEADERS = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 14_0) AppleWebKit/537.36 Chrome/128 Safari/537.36"}


@dataclass
class RentRange:
    municipality: str
    province: str
    ine_code: str
    year: int
    sample_size: int  # rented properties in the sample (collective housing)
    eur_m2_p25: float
    eur_m2_median: float
    eur_m2_p75: float


def download() -> Path:
    if not FILE.exists():
        FILE.parent.mkdir(parents=True, exist_ok=True)
        with httpx.stream("GET", URL, headers=HEADERS, timeout=300, follow_redirects=True) as response:
            response.raise_for_status()
            with FILE.open("wb") as file:
                for chunk in response.iter_bytes():
                    file.write(chunk)
    return FILE


def _key(name: str) -> str:
    return unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().lower().strip()


def _number(value: object) -> float | None:
    # Suppressed values come either as empty cells (None) or as empty strings ("").
    return None if value in (None, "") else float(value)


def load_municipalities(path: Path) -> dict[str, RentRange]:
    """Collective housing (VC) ranges of the last year, indexed by INE code (names are not unique: e.g. Cabanes)."""
    sheet = load_workbook(path, read_only=True)["Municipios"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    col = {name: i for i, name in enumerate(header)}
    ranges = {}
    for row in rows:
        median = _number(row[col[f"ALQM2_LV_M_VC_{YEAR}"]])
        if median is None:  # small municipalities are statistically suppressed
            continue
        ranges[row[col["CUMUN"]]] = RentRange(
            municipality=row[col["NMUN"]],
            province=row[col["NPRO"]],
            ine_code=row[col["CUMUN"]],
            year=2000 + int(YEAR),
            sample_size=int(_number(row[col[f"BI_ALVHEPCO_TVC_{YEAR}"]]) or 0),
            eur_m2_p25=round(_number(row[col[f"ALQM2_LV_25_VC_{YEAR}"]]), 2),
            eur_m2_median=round(median, 2),
            eur_m2_p75=round(_number(row[col[f"ALQM2_LV_75_VC_{YEAR}"]]), 2),
        )
    return ranges


def find_by_name(ranges: dict[str, RentRange], name: str) -> list[RentRange]:
    return [r for r in ranges.values() if _key(r.municipality) == _key(name)]


if __name__ == "__main__":
    municipality = sys.argv[1] if len(sys.argv) > 1 else "Madrid"
    surface = float(sys.argv[2]) if len(sys.argv) > 2 else 70

    ranges = load_municipalities(download())
    print(f"{len(ranges)} municipalities with {2000 + int(YEAR)} data")

    matches = find_by_name(ranges, municipality)
    if len(matches) != 1:
        sys.exit(f"{len(matches)} municipalities named {municipality!r}: {matches}")
    r = matches[0]
    print(r)
    print(
        f"Market range for {surface:g} m2 in {r.municipality}: "
        f"{r.eur_m2_p25 * surface:,.0f} - {r.eur_m2_p75 * surface:,.0f} EUR/month "
        f"(median {r.eur_m2_median * surface:,.0f})"
    )
